"""
Read a distributor price list (CSV or XLSX) into CatalogRow records, with
column mapping and a preview shown before anything is committed (spec
section 1).

CSV/XLSX only for now. PDF catalogs are not parsed here - no OCR/PDF
dependency is wired into this project yet, and guessing columns out of an
unstructured PDF table is exactly the kind of silent-inference risk this
feature is built to avoid. Point a PDF-only catalog at
pdf_conversion_instructions() for the manual route (save each page's table
as CSV/XLSX, e.g. via Excel's "From PDF" import or Acrobat's table export,
then import that) - rows are never silently dropped; a PDF input raises a
clear, actionable error instead.

COLUMN_ALIASES below includes real headers seen in catalogs/DC 55 July
Price List - copy.xlsx (a Keene Cheese Systems-style grocery distributor
export: ITEM/UPC/BRAND/DESCRIPTION/SIZE/UOM/PACK/WHOLESALE/RETAIL/...) plus
common synonyms. This is still a best-effort guess for any OTHER
distributor's file - once you import a real one, compare its header row
against this list and add any header that didn't get picked up. No other
code needs to change (same principle as smartscout_import.py).

IMPORTANT assumption this module cannot read from the file itself: whether
WHOLESALE-style price columns are priced "each" (per single sellable unit)
or "case"/"inner_pack" (per purchase unit). DC 55's own numbers only make
sense as per-each (e.g. $13.06 for an 11lb block of cheese sold in cases of
1 - consistent across wildly different pack sizes), so "each" is the
default here - but this is a GUESS, not a fact read from the file. It is
surfaced prominently in CatalogImportPreview.price_basis_note and must be
confirmed (or corrected) by a human before a real import is committed.
"""

import csv
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from catalog_models import (
    CatalogImportPreview, CatalogImportWarning, CatalogRow, ColumnMapping, PriceBasis,
)

# canonical_field -> list of header names (case-insensitive) that might
# appear in a real export. Add to these lists once you see your real file.
COLUMN_ALIASES: dict[str, list[str]] = {
    "supplier_sku": ["ITEM", "SKU", "Item #", "Item Number", "Supplier SKU", "Item Code"],
    "upc": ["UPC", "UPC/EAN", "EAN", "GTIN", "Barcode", "UPC Code", "UPC Number"],
    "brand": ["BRAND", "Brand", "Brand Name", "Manufacturer"],
    "description": ["DESCRIPTION", "Description", "Item Description", "Product Description", "Product Name", "Item Name"],
    "size": ["SIZE", "Size"],
    "uom": ["UOM", "Unit", "Unit of Measure"],
    "pack": ["PACK", "Pack", "Case Pack", "Units Per Case", "Case Qty", "Case Count"],
    "purchase_price": ["WHOLESALE", "Wholesale", "Cost", "Unit Cost", "Price", "Wholesale Price", "Our Price", "Net Price"],
    "country_of_origin": ["COUNTRY OF ORIGIN", "Country of Origin", "COO", "Origin"],
    "availability": ["ITEM STATUS", "Status", "Availability"],
    "distributor_sales_ytd_units": ["SYTD", "YTD Sold", "Sales YTD"],
    "distributor_trailing_12mo_units_sold": ["12MO SOLD", "12 Month Sold", "Trailing 12 Months Sold", "TTM Sold"],
    "category_1": ["CAT1", "Category", "Category 1"],
    "category_2": ["CAT2", "Category 2", "Subcategory"],
    "category_3": ["CAT3", "Category 3"],
}

REQUIRED_FIELDS = ["description", "purchase_price"]  # upc is highly recommended but not hard-required to preview

PDF_CONVERSION_NOTE = (
    "PDF catalog intake isn't implemented - no OCR/PDF extraction dependency is wired into this project yet. "
    "Manual route: open the PDF, export/save its product table as CSV or XLSX (Excel's Data > From PDF, or "
    "Acrobat's \"Export PDF\" > table/spreadsheet), then import that file instead. Review every extracted row "
    "against the original PDF page before treating it as final - OCR/table-extraction errors are common, "
    "especially on price and UPC columns."
)


def pdf_conversion_instructions() -> str:
    return PDF_CONVERSION_NOTE


def _build_header_lookup(columns: list[str]) -> dict[str, str]:
    lower_columns = {c.strip().lower(): c for c in columns}
    lookup: dict[str, str] = {}
    for canonical, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            match = lower_columns.get(alias.strip().lower())
            if match:
                lookup[canonical] = match
                break
    return lookup


def _cell_to_text(value) -> Optional[str]:
    """openpyxl hands back each cell's native type (str/int/float/None).
    A whole-number float (e.g. 12.0 for a PACK column Excel stored as a
    number) is rendered without the spurious ".0" so it matches what a
    human reading the sheet would type - everything else is just str()'d
    as-is. This never invents or strips a leading zero: that only happens
    for genuine text cells, which openpyxl already hands back as str."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _read_table(
    path: Path, max_rows: Optional[int] = None,
) -> tuple[list[str], list[dict[str, Optional[str]]]]:
    """Returns (header, rows) - rows as plain dicts of header -> text,
    uniformly for both CSV and XLSX. Deliberately NOT pandas: this
    project avoids it elsewhere (see smartscout_import.py's stdlib `csv`
    use) and pandas pulls in numpy, whose compiled extension has been seen
    to fail to load specifically under pytest's rootdir-resolved import
    path in this environment - openpyxl has no such dependency.

    max_rows caps how many DATA rows get materialized - openpyxl's own
    per-row XML parsing is the dominant cost on a large file (the DC 55
    sample is 98,719 rows and takes ~7s to read in full), so a fast
    preview of the first N rows must stop early rather than read
    everything and slice afterward."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        raise ValueError(PDF_CONVERSION_NOTE)
    if suffix in (".xlsx", ".xlsm"):
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            # Use the first non-empty sheet - a price list is rarely split
            # across sheets, but a trailing blank sheet (seen in the DC 55
            # sample file) is common and must not be silently preferred.
            for sheet_name in wb.sheetnames:
                rows_iter = wb[sheet_name].iter_rows(values_only=True)
                header_row = next(rows_iter, None)
                if header_row is None or all(v is None for v in header_row):
                    continue
                header = [_cell_to_text(h) or "" for h in header_row]
                records = []
                for raw_row in rows_iter:
                    if raw_row is None or all(v is None for v in raw_row):
                        continue
                    records.append({header[i]: _cell_to_text(v) for i, v in enumerate(raw_row) if i < len(header) and header[i]})
                    if max_rows is not None and len(records) >= max_rows:
                        break
                if records:
                    return header, records
            raise ValueError(f"No non-empty sheet found in {path}")
        finally:
            wb.close()
    if suffix in (".csv", ".tsv"):
        sep = "\t" if suffix == ".tsv" else ","
        with path.open(newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f, delimiter=sep)
            header = list(reader.fieldnames or [])
            records = []
            for r in reader:
                records.append(dict(r))
                if max_rows is not None and len(records) >= max_rows:
                    break
        return header, records
    if suffix == ".xls":
        raise ValueError("Legacy .xls isn't supported here - resave as .xlsx or .csv first.")
    raise ValueError(f"Unsupported catalog file type: {suffix} (supported: .csv, .tsv, .xlsx, .xlsm)")


def _count_rows(path: Path) -> int:
    """Cheap total-row count for CatalogImportPreview.total_row_count,
    without materializing any row content. For XLSX this reads the
    sheet's declared dimension (openpyxl metadata - no per-row XML
    parsing) rather than iterating every row: on the DC 55 sample file
    (98,719 rows) that's ~1s instead of ~7s. May slightly over-count if
    the sheet's used-range includes trailing blank rows."""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet_name in wb.sheetnames:
                max_row = wb[sheet_name].max_row
                if max_row and max_row > 1:
                    return max_row - 1  # minus the header row
            return 0
        finally:
            wb.close()
    if suffix in (".csv", ".tsv"):
        with path.open(newline="", encoding="utf-8-sig") as f:
            return max(sum(1 for _ in f) - 1, 0)
    return 0


def _scan_duplicate_upcs(path: Path, upc_header: str) -> int:
    """Full-file pass counting duplicate UPC occurrences. Deliberately
    separate from the fast preview path - this must visit every row, which
    is the expensive part on a large catalog (see _read_table's docstring)."""
    _, records = _read_table(path)
    seen: set[str] = set()
    duplicate_count = 0
    for row in records:
        upc = (row.get(upc_header) or "").strip()
        if not upc:
            continue
        if upc in seen:
            duplicate_count += 1
        seen.add(upc)
    return duplicate_count


def _clean(value) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return None
    return text


def _to_decimal(value) -> Optional[Decimal]:
    text = _clean(value)
    if text is None:
        return None
    text = text.replace("$", "").replace(",", "").strip()
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _to_int(value) -> Optional[int]:
    text = _clean(value)
    if text is None:
        return None
    try:
        return int(Decimal(text))
    except InvalidOperation:
        return None


_GTIN_VALID_LENGTHS = (8, 12, 13, 14)


def validate_gtin_checksum(code: str) -> Optional[bool]:
    """GS1 check-digit validation for an 8/12/13/14-digit UPC/EAN/GTIN.
    Returns None (not False) for anything that isn't purely digits of a
    recognized length - an unusual-but-possibly-valid supplier code is
    never flagged as definitively wrong, only as not checkable."""
    if not code or not code.isdigit() or len(code) not in _GTIN_VALID_LENGTHS:
        return None
    digits = [int(d) for d in code]
    *body, check_digit = digits
    total = 0
    # GS1 algorithm: working from the rightmost body digit, alternate
    # weights 3, 1, 3, 1, ...
    for i, digit in enumerate(reversed(body)):
        weight = 3 if i % 2 == 0 else 1
        total += digit * weight
    computed_check = (10 - (total % 10)) % 10
    return computed_check == check_digit


def suggest_leading_zero_correction(code: str) -> Optional[str]:
    """If `code` is shorter than a standard UPC-A length (12 digits) and
    zero-padding it up to 12 digits produces a checksum-valid UPC-A, return
    that reconstruction. This is a very common real-world artifact of a
    UPC column being treated as a number somewhere upstream (Excel, or a
    prior export) and losing its leading zero(s) - confirmed against the
    DC 55 sample catalog, where ~32% of UPCs are exactly one digit short of
    a checksum-valid UPC-A. Returns None when no such padding produces a
    valid checksum, so a genuinely different (non-truncated) code is never
    silently "fixed" into the wrong value. Never mutates the raw value -
    callers store this as a reviewable suggestion alongside it, not a
    replacement."""
    if not code or not code.isdigit() or len(code) >= 12:
        return None
    padded = code.zfill(12)
    return padded if validate_gtin_checksum(padded) else None


def _row_to_catalog_row(
    row: dict, lookup: dict[str, str], source_row: int, supplier_name: str,
    price_basis: PriceBasis, catalog_version: Optional[str], warnings: list[CatalogImportWarning],
) -> Optional[CatalogRow]:
    description = _clean(row.get(lookup.get("description", ""), None))
    if not description:
        warnings.append(CatalogImportWarning(source_row=source_row, message="No description/item name - row skipped"))
        return None

    upc = _clean(row.get(lookup.get("upc", ""), None))
    upc_valid = validate_gtin_checksum(upc) if upc else None
    upc_suggested_correction = None
    if upc and upc_valid is False:
        warnings.append(CatalogImportWarning(
            source_row=source_row, message=f"UPC '{upc}' fails the GS1 check-digit validation - kept as-is, flagged for review",
        ))
    elif upc and upc_valid is None:
        upc_suggested_correction = suggest_leading_zero_correction(upc)
        if upc_suggested_correction:
            warnings.append(CatalogImportWarning(
                source_row=source_row,
                message=f"UPC '{upc}' is {len(upc)} digits (not a standard length) but zero-pads to a "
                        f"checksum-valid UPC-A '{upc_suggested_correction}' - likely lost a leading zero upstream. "
                        "Kept as-is in `upc`; the reconstruction is in `upc_suggested_correction` for review.",
            ))

    purchase_price = _to_decimal(row.get(lookup.get("purchase_price", ""), None))
    if purchase_price is None:
        warnings.append(CatalogImportWarning(source_row=source_row, message="No purchase price found - row kept, cost will be unknown"))

    pack_value = _to_int(row.get(lookup.get("pack", ""), None)) or 1

    # See module docstring: PACK means something different depending on
    # the catalog's price basis. If the price is already "each", PACK is
    # the order multiple (must buy whole cases of this many); if the price
    # is per inner-pack/case, PACK is how many eaches make up that unit.
    if price_basis == "each":
        units_per_purchase_unit = 1
        order_multiple = pack_value
    else:
        units_per_purchase_unit = pack_value
        order_multiple = None

    categories = [
        c for c in (
            _clean(row.get(lookup.get("category_1", ""), None)),
            _clean(row.get(lookup.get("category_2", ""), None)),
            _clean(row.get(lookup.get("category_3", ""), None)),
        ) if c
    ]

    return CatalogRow(
        supplier_name=supplier_name,
        catalog_version=catalog_version,
        source_row=source_row,
        supplier_sku=_clean(row.get(lookup.get("supplier_sku", ""), None)),
        brand=_clean(row.get(lookup.get("brand", ""), None)),
        description=description,
        size=_clean(row.get(lookup.get("size", ""), None)),
        uom=_clean(row.get(lookup.get("uom", ""), None)),
        upc=upc,
        upc_valid=upc_valid,
        upc_suggested_correction=upc_suggested_correction,
        price_basis=price_basis,
        purchase_price=purchase_price,
        units_per_purchase_unit=units_per_purchase_unit,
        order_multiple=order_multiple,
        availability=_clean(row.get(lookup.get("availability", ""), None)),
        country_of_origin=_clean(row.get(lookup.get("country_of_origin", ""), None)),
        distributor_trailing_12mo_units_sold=_to_int(row.get(lookup.get("distributor_trailing_12mo_units_sold", ""), None)),
        distributor_sales_ytd_units=_to_int(row.get(lookup.get("distributor_sales_ytd_units", ""), None)),
        categories=categories,
        raw_fields={k: _clean(v) or "" for k, v in row.items()},
    )


def preview_catalog(
    path: str | Path,
    supplier_name: str,
    *,
    mapping: Optional[ColumnMapping] = None,
    sample_size: int = 25,
    catalog_version: Optional[str] = None,
    thorough: bool = False,
) -> CatalogImportPreview:
    """Detect columns, build a handful of sample rows, and surface the
    price-basis assumption - all with NO commitment to any store. Call
    load_catalog_rows() separately once a human has reviewed this.

    Fast by default: only reads the first sample_size rows plus a cheap
    row-count (see _count_rows) - on the 98,719-row DC 55 sample file
    that's ~1s instead of ~8s. Pass thorough=True to also run the
    full-file duplicate-UPC scan before a real commit (duplicate_upc_count
    stays None, meaning "not checked," until you do)."""
    path = Path(path)
    columns, records = _read_table(path, max_rows=sample_size)
    lookup = mapping.mapping if mapping else _build_header_lookup(columns)
    price_basis = mapping.price_basis if mapping else "each"

    unmapped_required = [f for f in REQUIRED_FIELDS if f not in lookup]

    warnings: list[CatalogImportWarning] = []
    sample_rows: list[CatalogRow] = []
    for i, row in enumerate(records, start=2):  # row 1 is the header
        catalog_row = _row_to_catalog_row(row, lookup, i, supplier_name, price_basis, catalog_version, warnings)
        if catalog_row:
            sample_rows.append(catalog_row)

    duplicate_upc_count = None
    upc_header = lookup.get("upc")
    if thorough and upc_header:
        duplicate_upc_count = _scan_duplicate_upcs(path, upc_header)

    return CatalogImportPreview(
        detected_mapping=lookup,
        unmapped_required_fields=unmapped_required,
        price_basis_assumption=price_basis,
        price_basis_note=(
            f'ASSUMPTION, not read from the file: purchase-price column is treated as price-per-"{price_basis}". '
            "Confirm this against the supplier's actual price sheet before committing a real import - "
            "see catalog_import.py's module docstring for why this can't be auto-detected."
        ),
        sample_rows=sample_rows,
        total_row_count=_count_rows(path),
        warnings=warnings,
        duplicate_upc_count=duplicate_upc_count,
    )


def load_catalog_rows(
    path: str | Path,
    supplier_name: str,
    *,
    mapping: Optional[ColumnMapping] = None,
    catalog_version: Optional[str] = None,
    limit: Optional[int] = None,
) -> tuple[list[CatalogRow], list[CatalogImportWarning]]:
    """Full import, for real (after preview_catalog has been reviewed).
    Returns (rows, warnings) - rows missing a usable description are
    skipped (with a warning), never silently dropped without a trace."""
    path = Path(path)
    columns, records = _read_table(path, max_rows=limit)
    lookup = mapping.mapping if mapping else _build_header_lookup(columns)
    price_basis = mapping.price_basis if mapping else "each"

    rows: list[CatalogRow] = []
    warnings: list[CatalogImportWarning] = []
    for i, row in enumerate(records, start=2):
        catalog_row = _row_to_catalog_row(row, lookup, i, supplier_name, price_basis, catalog_version, warnings)
        if catalog_row:
            rows.append(catalog_row)

    return rows, warnings
